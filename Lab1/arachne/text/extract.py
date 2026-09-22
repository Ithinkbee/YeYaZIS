"""Извлечение текста и заголовка из файлов разных форматов.

Поддерживаются форматы, типичные для файловых ресурсов ЛВС:
txt, md, csv, html, docx, pdf, rtf.
"""

from __future__ import annotations

import re
from pathlib import Path

ENCODINGS = ("utf-8-sig", "utf-8", "cp1251", "koi8-r", "cp866")


class ExtractionError(Exception):
    """Файл не удалось разобрать."""


def _decode(raw: bytes) -> str:
    """Подбирает кодировку: сначала utf-8, затем кириллические однобайтовые."""
    best_text, best_score = "", -1.0
    for enc in ENCODINGS:
        try:
            text = raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        # оценка качества: доля кириллицы/латиницы/пробелов среди символов
        good = sum(1 for ch in text[:4000] if ch.isalnum() or ch.isspace() or ch in ".,;:!?-—«»()")
        score = good / max(1, len(text[:4000]))
        if enc.startswith("utf-8") and score > 0.85:
            return text
        if score > best_score:
            best_text, best_score = text, score
    if not best_text:
        raise ExtractionError("не удалось определить кодировку файла")
    return best_text


def _clean(text: str) -> str:
    """Схлопывает пробелы и пустые строки."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _title_from_text(text: str, fallback: str) -> str:
    """Заголовок = первая содержательная строка, если она похожа на заголовок."""
    for line in text.split("\n"):
        line = line.strip(" #*=-\t")
        if 3 <= len(line) <= 120:
            return line
        if line:
            break
    return fallback


# --- Обработчики форматов ---------------------------------------------------

def _extract_plain(path: Path) -> tuple[str, str]:
    text = _clean(_decode(path.read_bytes()))
    return _title_from_text(text, path.stem), text


def _extract_html(path: Path) -> tuple[str, str]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(path.read_bytes()), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    if not title:
        h1 = soup.find(["h1", "h2"])
        if h1:
            title = h1.get_text(" ", strip=True)
    text = _clean(soup.get_text("\n", strip=True))
    return title or path.stem, text


def _extract_docx(path: Path) -> tuple[str, str]:
    import docx  # python-docx

    document = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    title = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            if paragraph.style is not None and "Heading" in (paragraph.style.name or ""):
                title = paragraph.text.strip()
            break
    text = _clean("\n".join(p for p in paragraphs if p))
    core_title = ""
    try:
        core_title = (document.core_properties.title or "").strip()
    except Exception:
        pass
    return title or core_title or _title_from_text(text, path.stem), text


def _extract_pdf(path: Path) -> tuple[str, str]:
    import pypdf

    reader = pypdf.PdfReader(str(path))
    chunks = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    text = _clean("\n".join(chunks))
    title = ""
    try:
        title = (reader.metadata.title or "").strip() if reader.metadata else ""
    except Exception:
        title = ""
    return title or _title_from_text(text, path.stem), text


def _extract_rtf(path: Path) -> tuple[str, str]:
    raw = _decode(path.read_bytes())
    try:
        from striprtf.striprtf import rtf_to_text

        text = rtf_to_text(raw, errors="ignore")
    except ImportError:
        # запасной разбор: убираем управляющие последовательности RTF
        text = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), raw)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
        text = text.replace("{", " ").replace("}", " ")
    text = _clean(text)
    return _title_from_text(text, path.stem), text


EXTRACTORS = {
    ".txt": _extract_plain,
    ".md": _extract_plain,
    ".csv": _extract_plain,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
    ".rtf": _extract_rtf,
}


def supported(path: Path) -> bool:
    return path.suffix.lower() in EXTRACTORS


def extract(path: Path) -> tuple[str, str]:
    """Возвращает (заголовок, текст) документа.

    Бросает ExtractionError, если формат не поддержан или файл повреждён.
    """
    handler = EXTRACTORS.get(path.suffix.lower())
    if handler is None:
        raise ExtractionError(f"формат {path.suffix} не поддерживается")
    try:
        title, text = handler(path)
    except ExtractionError:
        raise
    except Exception as exc:  # повреждённый файл не должен ронять паука
        raise ExtractionError(f"{type(exc).__name__}: {exc}") from exc
    if not text.strip():
        raise ExtractionError("документ не содержит текста")
    return title.strip() or path.stem, text
