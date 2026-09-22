"""Генератор тестовой коллекции документов и эталонной разметки.

Раскладывает документы по каталогам, изображающим узлы локальной сети
(NODE-A, NODE-B, NODE-C), в разных форматах (txt, md, html, docx, pdf),
и сохраняет эталонный набор запросов с разметкой релевантности.

Запуск:  python tools/make_collection.py
"""

from __future__ import annotations

import argparse
import csv
import html
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from arachne import config  # noqa: E402
from tools.collection_data import EVAL_QUERIES, TOPICS  # noqa: E402

NODES = ["NODE-A", "NODE-B", "NODE-C"]
FORMATS = [".txt", ".html", ".docx", ".pdf", ".md"]


def normalize(text: str) -> str:
    """Убирает переносы строк, оставшиеся от форматирования исходника."""
    return re.sub(r"\s+", " ", text).strip()


def paragraphs(text: str, count: int = 2) -> list[str]:
    """Делит текст на несколько абзацев по границам предложений."""
    sentences = re.findall(r"[^.!?]+[.!?]+", normalize(text)) or [normalize(text)]
    if count <= 1 or len(sentences) < 4:
        return [" ".join(s.strip() for s in sentences)]
    middle = (len(sentences) + 1) // 2
    return [
        " ".join(s.strip() for s in sentences[:middle]),
        " ".join(s.strip() for s in sentences[middle:]),
    ]


# --- Запись форматов --------------------------------------------------------

def write_txt(path: Path, title: str, text: str) -> None:
    body = "\n\n".join(paragraphs(text))
    path.write_text(f"{title}\n\n{body}\n", encoding="utf-8")


def write_md(path: Path, title: str, text: str) -> None:
    body = "\n\n".join(paragraphs(text))
    path.write_text(f"# {title}\n\n{body}\n", encoding="utf-8")


def write_html(path: Path, title: str, text: str) -> None:
    body = "\n".join(f"  <p>{html.escape(p)}</p>" for p in paragraphs(text))
    path.write_text(
        "<!DOCTYPE html>\n<html lang=\"ru\">\n<head>\n"
        "  <meta charset=\"utf-8\">\n"
        f"  <title>{html.escape(title)}</title>\n</head>\n<body>\n"
        f"  <h1>{html.escape(title)}</h1>\n{body}\n</body>\n</html>\n",
        encoding="utf-8",
    )


def write_docx(path: Path, title: str, text: str) -> None:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    for paragraph in paragraphs(text):
        document.add_paragraph(paragraph)
    document.save(str(path))


def write_pdf(path: Path, title: str, text: str) -> None:
    import textwrap

    import matplotlib

    matplotlib.use("Agg")
    matplotlib.rcParams["pdf.fonttype"] = 42  # TrueType: текст извлекается корректно
    import matplotlib.pyplot as plt

    figure = plt.figure(figsize=(8.27, 11.69))  # A4
    figure.text(0.08, 0.94, title, fontsize=15, va="top", wrap=True)
    y = 0.88
    for paragraph in paragraphs(text):
        for line in textwrap.wrap(paragraph, width=88):
            figure.text(0.08, y, line, fontsize=10, va="top")
            y -= 0.021
        y -= 0.015
    figure.savefig(str(path))
    plt.close(figure)


WRITERS = {
    ".txt": write_txt,
    ".md": write_md,
    ".html": write_html,
    ".docx": write_docx,
    ".pdf": write_pdf,
}


# --- Сборка коллекции -------------------------------------------------------

def build(clean: bool = True) -> dict:
    root = config.COLLECTION_DIR
    if clean and root.exists():
        shutil.rmtree(root)
    root.mkdir(parents=True, exist_ok=True)

    created: list[tuple[str, str, str]] = []
    position = 0
    for topic_key, topic in TOPICS.items():
        for document in topic["docs"]:
            node = NODES[position % len(NODES)]
            extension = FORMATS[position % len(FORMATS)]
            directory = root / node / topic_key
            directory.mkdir(parents=True, exist_ok=True)
            path = directory / f"{document['slug']}{extension}"
            WRITERS[extension](path, document["title"], document["text"])
            created.append((document["slug"], node, extension))
            position += 1

    return {
        "documents": len(created),
        "nodes": NODES,
        "topics": {key: len(value["docs"]) for key, value in TOPICS.items()},
        "formats": {
            extension: sum(1 for _, _, e in created if e == extension) for extension in FORMATS
        },
    }


def known_slugs() -> set[str]:
    """Идентификаторы всех документов коллекции."""
    return {document["slug"] for topic in TOPICS.values() for document in topic["docs"]}


def check_consistency() -> None:
    """Проверяет, что коллекция и разметка согласованы.

    Опечатка в идентификаторе документа внутри qrels молча искажает все метрики
    качества: несуществующий документ просто никогда не находится. Поэтому
    несоответствие останавливает генерацию.
    """
    slugs = [document["slug"] for topic in TOPICS.values() for document in topic["docs"]]
    duplicates = sorted({slug for slug in slugs if slugs.count(slug) > 1})
    if duplicates:
        raise SystemExit(f"повторяющиеся идентификаторы документов: {', '.join(duplicates)}")

    known = set(slugs)
    unknown: set[str] = set()
    for query in EVAL_QUERIES:
        unknown |= set(query["relevant"]) - known
        unknown |= set(query.get("irrelevant", [])) - known
    if unknown:
        raise SystemExit(
            "разметка ссылается на несуществующие документы: " + ", ".join(sorted(unknown))
        )

    without_relevant = [query["text"] for query in EVAL_QUERIES if not query["relevant"]]
    if without_relevant:
        raise SystemExit("запросы без единого релевантного документа: "
                         + "; ".join(without_relevant))


def write_eval_set() -> tuple[int, int]:
    """Сохраняет эталонные запросы и разметку релевантности."""
    config.ensure_dirs()
    queries_path = config.DATA_DIR / "eval_queries.csv"
    with queries_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["query_id", "text", "note"])
        for number, query in enumerate(EVAL_QUERIES, 1):
            writer.writerow([number, query["text"], query.get("note", "")])

    pairs = 0
    with config.QRELS_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["query_id", "doc_slug", "rel"])
        for number, query in enumerate(EVAL_QUERIES, 1):
            judgements = dict(query["relevant"])
            for slug in query.get("irrelevant", []):
                judgements.setdefault(slug, 0)
            for slug, rel in sorted(judgements.items()):
                writer.writerow([number, slug, rel])
                pairs += 1
    return len(EVAL_QUERIES), pairs


def main() -> None:
    parser = argparse.ArgumentParser(description="Генерация тестовой коллекции ИПС «Арахна»")
    parser.add_argument(
        "--keep", action="store_true", help="не удалять существующие файлы коллекции"
    )
    arguments = parser.parse_args()

    check_consistency()
    summary = build(clean=not arguments.keep)
    queries, pairs = write_eval_set()

    print(f"Коллекция: {summary['documents']} документов в {config.COLLECTION_DIR}")
    print("  узлы:    " + ", ".join(summary["nodes"]))
    print("  темы:    " + ", ".join(f"{k}={v}" for k, v in summary["topics"].items()))
    print("  форматы: " + ", ".join(f"{k}={v}" for k, v in summary["formats"].items()))
    print(f"Эталон: {queries} запросов, {pairs} оценок релевантности")


if __name__ == "__main__":
    main()
