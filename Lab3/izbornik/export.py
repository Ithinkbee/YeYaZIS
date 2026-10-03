"""Сохранение реферата в файл: TXT, HTML, DOCX, JSON, SCs.

Каждый формат содержит оба раздела реферата — классический реферат и реферат
в виде списка ключевых слов — и ссылку на исходный документ.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import datetime

from izbornik import APP_NAME, VERSION, config
from izbornik.summary import Summary

FORMATS: dict[str, tuple[str, str]] = {
    "txt": ("Текст (TXT)", "text/plain; charset=utf-8"),
    "html": ("Веб-страница (HTML)", "text/html; charset=utf-8"),
    "docx": ("Документ Word (DOCX)", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "json": ("Данные (JSON)", "application/json; charset=utf-8"),
    "scs": ("Фрагмент базы знаний (SCs)", "text/plain; charset=utf-8"),
}


@dataclass
class Context:
    """Сведения о документе, которых нет в самом реферате."""

    source_link: str = ""        # адрес исходного документа в системе
    source_site: str = ""        # сайт-источник
    license: str = ""
    authors: str = ""
    compressed: bool = True      # выводить сжатые предложения


def filename(summary: Summary, fmt: str) -> str:
    base = summary.doc_id or "text"
    base = re.sub(r"[^\w\-]+", "_", base)
    return f"referat-{base}-{summary.requested}.{fmt}"


def _sentence(summary: Summary, index: int, ctx: Context) -> str:
    sentence = summary.sentences[index]
    return sentence.compressed if ctx.compressed else sentence.text


def _engine(summary: Summary) -> str:
    if summary.engine == "ostis":
        return "sc-агент OSTIS (реферат прочитан из базы знаний)"
    return "локальный расчёт"


# --- TXT ------------------------------------------------------------------------

def to_txt(summary: Summary, ctx: Context) -> str:
    lines = [
        f"РЕФЕРАТ ДОКУМЕНТА «{summary.title}»",
        "",
        f"Исходный документ: {ctx.source_link or '—'}",
    ]
    if summary.source_url:
        lines.append(f"Источник: {ctx.source_site} {summary.source_url}".strip())
    if ctx.authors:
        lines.append(f"Авторы: {ctx.authors}")
    if ctx.license:
        lines.append(f"Лицензия источника: {ctx.license}")
    lines += [
        f"Язык: {config.language_name(summary.language)}"
        + (f"; предметная область: {config.domain_name(summary.domain)}" if summary.domain else ""),
        f"Метод: sentence extraction; построено: {_engine(summary)}",
        "",
        f"1. КЛАССИЧЕСКИЙ РЕФЕРАТ ({len(summary.sentences)} предложений)",
        "",
    ]
    for i, s in enumerate(summary.sentences):
        lines.append(f"  {i + 1:2}. {_sentence(summary, i, ctx)}")
        lines.append(f"      [предложение № {s.index + 1}; вес {s.weight:.3f} = "
                     f"Score {s.score:.3f} × Posd {s.posd:.3f} × Posp {s.posp:.3f}; место {s.rank}]")
    lines += ["", "2. РЕФЕРАТ В ВИДЕ СПИСКА КЛЮЧЕВЫХ СЛОВ", ""]
    for top in summary.keywords.tree:
        for depth, kw in top.walk():
            lines.append("    " * (depth + 1) + kw.text)
    if summary.keywords.loose:
        lines += ["", "  Другие словосочетания: " + ", ".join(k.text for k in summary.keywords.loose)]
    stats = summary.stats
    lines += [
        "",
        f"Объём документа: {stats.get('chars', 0)} символов; реферата: {stats.get('summary_chars', 0)} "
        f"символов (коэффициент сжатия {stats.get('compression', 0):.3f}).",
        f"«{APP_NAME}» {VERSION}, {datetime.now():%d.%m.%Y %H:%M}.",
    ]
    return "\n".join(lines) + "\n"


# --- HTML -----------------------------------------------------------------------

def to_html(summary: Summary, ctx: Context, render) -> str:
    """HTML-страница; render — функция шаблонизатора веб-интерфейса."""
    return render("export.html", summary=summary, ctx=ctx, engine=_engine(summary),
                  generated=datetime.now().strftime("%d.%m.%Y %H:%M"))


# --- DOCX -----------------------------------------------------------------------

def to_docx(summary: Summary, ctx: Context) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = docx.Document()
    section = document.sections[0]
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def para(text: str = "", *, bold=False, size=None, align=None, italic=False, color=None, indent=None):
        p = document.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.left_indent = indent
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        return p

    para(f"Реферат документа «{summary.title}»", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    meta = [f"Исходный документ: {ctx.source_link}"]
    if summary.source_url:
        meta.append(f"Источник: {ctx.source_site} {summary.source_url}")
    if ctx.authors:
        meta.append(f"Авторы: {ctx.authors}")
    meta.append(f"Язык: {config.language_name(summary.language)}"
                + (f"; область: {config.domain_name(summary.domain)}" if summary.domain else ""))
    meta.append(f"Метод: sentence extraction; построено: {_engine(summary)}")
    for line in meta:
        para(line, size=11, color=(0x55, 0x55, 0x55))

    para(f"1. Классический реферат ({len(summary.sentences)} предложений)", bold=True)
    for i, s in enumerate(summary.sentences):
        p = para(f"{i + 1}. {_sentence(summary, i, ctx)}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.first_line_indent = Cm(1.25)
        para(f"предложение № {s.index + 1}; вес {s.weight:.3f} (Score {s.score:.3f} × "
             f"Posd {s.posd:.3f} × Posp {s.posp:.3f}); место по весу {s.rank}",
             size=10, italic=True, color=(0x77, 0x77, 0x77), indent=Cm(1.25))

    para("2. Реферат в виде списка ключевых слов", bold=True)
    for top in summary.keywords.tree:
        for depth, kw in top.walk():
            para(kw.text, bold=depth == 0, indent=Cm(1.0 + depth * 1.0))
    if summary.keywords.loose:
        para("Другие словосочетания: " + ", ".join(k.text for k in summary.keywords.loose), size=12)

    stats = summary.stats
    para(f"Объём документа — {stats.get('chars', 0)} символов, реферата — {stats.get('summary_chars', 0)} "
         f"(коэффициент сжатия {stats.get('compression', 0):.3f}). «{APP_NAME}» {VERSION}, "
         f"{datetime.now():%d.%m.%Y %H:%M}.", size=10, color=(0x77, 0x77, 0x77))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- JSON и SCs --------------------------------------------------------------------

def to_json(summary: Summary, ctx: Context) -> str:
    data = summary.to_dict()
    data["source_link"] = ctx.source_link
    data["generator"] = f"{APP_NAME} {VERSION}"
    return json.dumps(data, ensure_ascii=False, indent=2)


def to_scs(summary: Summary, ctx: Context) -> str:
    from izbornik.ostis.scs import to_scs as scs

    return scs(summary)


def render(summary: Summary, fmt: str, ctx: Context, template_render=None) -> bytes:
    if fmt == "txt":
        return to_txt(summary, ctx).encode("utf-8")
    if fmt == "html":
        return to_html(summary, ctx, template_render).encode("utf-8")
    if fmt == "docx":
        return to_docx(summary, ctx)
    if fmt == "json":
        return to_json(summary, ctx).encode("utf-8")
    if fmt == "scs":
        return to_scs(summary, ctx).encode("utf-8")
    raise ValueError(f"неизвестный формат: {fmt}")
